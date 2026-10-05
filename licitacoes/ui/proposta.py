"""Telas reais do Passo 1 (upload do edital) e Passo 2 (proposta editável).

Sem dados de exemplo: tudo vem do edital enviado e do banco SQLite.
Uso no app.py:
    from licitacoes.ui.proposta import render_edital_page, render_proposta_page
    def page_edital():   render_edital_page(on_done=lambda: ir_para("proposta"))
    def page_proposta(): render_proposta_page()
"""
from __future__ import annotations

import re
import io
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from licitacoes.config import DATA_DIR
from licitacoes.edital.items_parser import parse_edital
from licitacoes.services import persistence as db
from licitacoes.services.money import format_brl, format_plain, parse_brl_to_cents
from licitacoes.services.proposal_calc import status, summarize
from licitacoes.services.company import CompanyService

EDITAIS_DIR = DATA_DIR / "editais"


def _safe_name(name: str) -> str:
    return re.sub(r"[^\w.\- ]", "_", name)


# ------------------------------------------------------------------ Passo 1
def render_edital_page(on_done=None, db_path=None):
    st.header("1. Edital")
    st.write("Envie o edital (arquivo .md, .docx ou .pdf). Eu leio a lista de itens e os preços máximos.")
    up = st.file_uploader("Arraste o edital aqui", type=["md", "docx", "pdf", "txt"], key="edital_upload")

    if up is not None and st.button("Analisar edital", type="primary"):
        EDITAIS_DIR.mkdir(parents=True, exist_ok=True)
        dest = EDITAIS_DIR / _safe_name(up.name)
        dest.write_bytes(up.getvalue())
        with st.spinner("Lendo o edital..."):
            tid, res, ja_existia = import_edital_file(dest, db_path)
        if not res.items:
            for w in res.warnings:
                st.error(w)
            st.warning("Nada foi inventado: nenhum item foi criado. Você pode preencher os itens à mão.")
            return
        st.session_state["tender_id"] = tid
        if ja_existia:
            st.info("Este edital já foi analisado. Abri a proposta que você já tinha.")
        else:
            for w in res.warnings:
                st.warning(w)
            st.success(f"Li {len(res.items)} itens. Total pelos preços máximos: {format_brl(res.computed_total_cents)}.")
        if on_done:
            on_done()


def import_edital_file(path, db_path=None):
    """Lê o edital, grava no banco e devolve (tender_id | None, resultado, já_existia)."""
    path = Path(path)
    res = parse_edital(path)
    if not res.items:
        return None, res, False
    existing = db.find_tender_by_sha(res.sha256, db_path)
    if existing:
        return existing["id"], res, True
    tid = db.create_tender(
        path.stem, res.items, estimated_total_cents=res.estimated_total_cents,
        edital_sha256=res.sha256, edital_path=path, extraction_method=res.method,
        needs_review=res.needs_review, db_path=db_path)
    return tid, res, False


# ------------------------------------------------------------------ Passo 2
def _current_tender(db_path=None):
    tid = st.session_state.get("tender_id")
    t = db.get_tender(tid, db_path) if tid else None
    if t is None:
        abertas = db.list_tenders(db_path=db_path)
        t = abertas[0] if abertas else None
        if t:
            st.session_state["tender_id"] = t["id"]
    return t


def render_proposta_page(db_path=None):
    st.header("2. Minha proposta")
    t = _current_tender(db_path)
    if t is None:
        st.info("Ainda não há edital analisado. Volte ao Passo 1 e envie o edital.")
        return
    items = db.get_items(t["id"], db_path)
    if not items:
        st.warning("Esta licitação não tem itens.")
        return

    st.caption(f"Licitação: **{t['name']}**")
    if t["needs_review"]:
        st.warning("Atenção: a leitura do edital pede conferência. Compare os itens com o edital.")

    df = pd.DataFrame({
        "Item": [i["number"] for i in items],
        "Especificações": [i["description"] for i in items],
        "Unid.": [i["unit"] for i in items],
        "Qtde": [i["quantity"] for i in items],
        "Preço máximo": [format_brl(i["ceiling_cents"]) for i in items],
        "Marca": [i["brand"] for i in items],
        "Meu preço unitário (R$)": [format_plain(i["proposed_cents"]) for i in items],
        "Meu total": [format_brl(i["quantity"] * i["proposed_cents"]) if i["proposed_cents"] is not None else "" for i in items],
        "Situação": [_label(status(i["ceiling_cents"], i["proposed_cents"])) for i in items],
    })
    st.caption("Preencha **Marca** e **Meu preço unitário** (ex.: 8,50). Salva sozinho.")
    edited = st.data_editor(
        df, hide_index=True, use_container_width=True, key=f"proposta_{t['id']}",
        disabled=["Item", "Especificações", "Unid.", "Qtde", "Preço máximo", "Meu total", "Situação"],
        column_config={
            "Especificações": st.column_config.TextColumn(width="large"),
            "Meu preço unitário (R$)": st.column_config.TextColumn(help="Digite com vírgula, ex.: 8,50"),
        })

    # --- grava o que mudou (texto -> centavos, sem float)
    invalid, changes = [], []
    for it, (_, row) in zip(items, edited.iterrows()):
        txt = str(row["Meu preço unitário (R$)"] or "").strip()
        cents = parse_brl_to_cents(txt) if txt else None
        if txt and cents is None:
            invalid.append(it["number"])
            cents = it["proposed_cents"]
        brand = str(row["Marca"] or "").strip()
        if cents != it["proposed_cents"] or brand != it["brand"].strip():
            changes.append((it["number"], cents, brand))
    if invalid:
        st.error("Preço inválido nos itens: " + ", ".join(map(str, invalid)) + ". Use o formato 8,50.")
    if changes:
        db.save_proposal_bulk(t["id"], changes, db_path)
        st.rerun()

    s = summarize(items)
    c1, c2 = st.columns(2)
    c1.metric("Total da minha proposta", format_brl(s["total_proposed_cents"]))
    c2.metric("Total pelos preços máximos", format_brl(s["total_ceiling_cents"]))
    if s["above_ceiling"]:
        st.error("🔴 Acima do preço máximo (será DESCLASSIFICADO) nos itens: "
                 + ", ".join(map(str, s["above_ceiling"])))
    if s["missing_price"]:
        st.warning("Faltam preços nos itens: " + ", ".join(map(str, s["missing_price"])))
    if s["missing_brand"]:
        st.warning("Faltam marcas nos itens: " + ", ".join(map(str, s["missing_brand"])))

    # --- Exportação para Excel ---
    st.divider()
    col_exp1, col_exp2 = st.columns([1, 3])
    with col_exp1:
        if st.button("💾 Exportar Proposta (Excel)", type="primary"):
            company = CompanyService.get_company_data()
            excel_data = _export_to_excel(t, company, items)
            st.download_button(
                label="⬇️ Baixar Arquivo .xlsx",
                data=excel_data,
                file_name=f"Proposta_{_safe_name(t['name'])}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )


def _label(s: str) -> str:
    return {"ok": "✔ OK", "falta": "— sem preço", "acima": "🔴 Acima do máximo: desclassificado"}[s]


def _export_to_excel(tender, company, items):
    """Gera um arquivo Excel com cabeçalho e a tabela de proposta."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Proposta"

    # Estilos
    bold_font = Font(bold=True)
    center_align = Alignment(horizontal="center")

    # --- Cabeçalho ---
    # Info da Empresa
    ws["A1"] = "PROPOSTA COMERCIAL"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A3"] = "EMPRESA:"
    ws["A3"].font = bold_font
    ws["B3"] = f"{company.razao_social} | CNPJ: {company.cnpj}"

    ws["A4"] = "ENDEREÇO:"
    ws["A4"].font = bold_font
    ws["B4"] = company.endereco

    ws["A5"] = "CONTATO:"
    ws["A5"].font = bold_font
    ws["B5"] = f"{company.email} | {company.telefone}"

    ws.merge_cells("B3:E3")
    ws.merge_cells("B4:E4")
    ws.merge_cells("B5:E5")

    # Info da Licitação
    ws["A7"] = "LICITAÇÃO:"
    ws["A7"].font = bold_font
    ws["B7"] = tender["name"]

    ws["A8"] = "ÓRGÃO:"
    ws["A8"].font = bold_font
    ws["B8"] = tender.get("organ", "Não informado")

    ws["A9"] = "PROCESSO:"
    ws["A9"].font = bold_font
    ws["B9"] = tender.get("process_number", "Não informado")

    ws.merge_cells("B7:E7")
    ws.merge_cells("B8:E8")
    ws.merge_cells("B9:E9")

    # --- Tabela de Itens ---
    headers = ["Item", "Especificações", "Unid.", "Qtde", "Marca", "Preço Unit. (R$)", "Total (R$)"]
    ws.append([]) # Linha em branco
    ws.append(headers)

    # Estilizar cabeçalho da tabela
    for cell in ws[ws.max_row]:
        cell.font = bold_font
        cell.alignment = center_align

    total_proposal = 0
    for i in items:
        unit_price = i["proposed_cents"] or 0
        total_item = i["quantity"] * unit_price
        total_proposal += total_item

        ws.append([
            i["number"],
            i["description"],
            i["unit"],
            i["quantity"],
            i["brand"],
            unit_price / 100,
            total_item / 100
        ])

    # Total Final
    ws.append([])
    ws.append(["", "", "", "", "TOTAL GERAL:", "", total_proposal / 100])
    ws.cell(row=ws.max_row, column=5).font = bold_font
    ws.cell(row=ws.max_row, column=6).font = bold_font

    # Formatação de colunas de preço
    for row in ws.iter_rows(min_row=12, max_row=ws.max_row, min_col=6, max_col=7):
        for cell in row:
            cell.number_format = '#,##0.00'

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
