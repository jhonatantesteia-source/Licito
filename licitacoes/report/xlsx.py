from typing import List
from pathlib import Path
from datetime import date
import json
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from licitacoes.config import settings
from licitacoes.edital.schema import EditalSchema

class ProposalGenerator:
    """Generates an Excel proposal based on the EditalSchema."""

    def generate(self, edital: EditalSchema, output_path: Path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Proposta de Preços"

        # Header
        ws.merge_cells('A1:E1')
        ws['A1'] = f"{edital.organ} - {edital.modality} {edital.process_number}"
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')

        # Company Info
        company = settings.company
        ws['A3'] = "Razão Social:"
        ws['B3'] = company.razao_social
        ws['A4'] = "CNPJ:"
        ws['B4'] = company.cnpj
        ws['A5'] = "Endereço:"
        ws['B5'] = company.endereco
        ws['A6'] = "Telefone/Email:"
        ws['B6'] = f"{company.telefone} / {company.email}"

        # Table Header
        headers = ["Item", "Descrição", "Unidade", "Quantidade", "Preço Unit. (R$)", "Preço Total (R$)", "Marca"]
        ws.append([]) # Blank line
        ws.append(headers)

        # Style headers
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')

        # Items
        for idx, item in enumerate(edital.items, 1):
            # We use formulas for totals: Total = Qty * Unit Price
            # Assuming Col D = Qty, Col E = Unit Price, Col F = Total
            row_num = ws.max_row + 1
            ws.append([
                item.id,
                item.description,
                item.unit,
                item.quantity,
                "", # Unit price to be filled by user
                f"=D{row_num}*E{row_num}",
                ""  # Brand to be filled by user
            ])

        # Grand Total
        last_row = ws.max_row
        total_row = last_row + 1
        ws.append([])
        ws.append(["", "", "", "", "TOTAL GERAL:", f"=SUM(F{ws.max_row-len(edital.items)+1}:F{last_row})", ""])
        ws.cell(row=total_row+1, column=5).font = Font(bold=True)
        ws.cell(row=total_row+1, column=6).font = Font(bold=True)

        wb.save(output_path)

# Singleton
proposal_gen = ProposalGenerator()
