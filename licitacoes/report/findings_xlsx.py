from typing import List, Dict, Any
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from licitacoes.rules.base import Finding

class FindingsReportGenerator:
    """Generates a professional Excel report of all findings for all bidders."""

    def generate(self, all_bidders_findings: Dict[str, List[Finding]], output_path: Path):
        wb = openpyxl.Workbook()

        # 1. Summary Sheet
        ws_summary = wb.active
        ws_summary.title = "Resumo Geral"

        headers = ["Licitante", "Achados Críticos", "Achados Altos", "Achados Médios", "Status Final"]
        ws_summary.append(headers)

        # Style Summary Headers
        for cell in ws_summary[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')

        # 2. Individual Bidder Sheets
        for bidder, findings in all_bidders_findings.items():
            # Add to summary
            critical = len([f for f in findings if f.severity == "critical"])
            high = len([f for f in findings if f.severity == "high"])
            med = len([f for f in findings if f.severity == "medium"])
            ws_summary.append([bidder, critical, high, med, "Análise Pendente"])

            # Create detail sheet
            ws_detail = wb.create_sheet(title=bidder[:30]) # Excel limit 31 chars

            detail_headers = ["Documento", "Achado", "Gravidade", "Regra", "Evidência", "Arquivo/Pág", "Ação Sugerida", "Status"]
            ws_detail.append(detail_headers)

            for cell in ws_detail[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal='center')

            for f in findings:
                ws_detail.append([
                    f.document,
                    f.finding_text,
                    f.severity.upper(),
                    f.rule_violated,
                    f.evidence,
                    f"{f.file_path} (p.{f.page or 'N/A'})",
                    f.suggested_action,
                    f.status
                ])

                # Color code severity
                last_row = ws_detail.max_row
                severity_cell = ws_detail.cell(row=last_row, column=3)
                if f.severity == "critical":
                    severity_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                elif f.severity == "high":
                    severity_cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

        # Auto-adjust columns
        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except: pass
                sheet.column_dimensions[column].width = max_length + 2

        wb.save(output_path)

# Singleton
findings_report_gen = FindingsReportGenerator()
