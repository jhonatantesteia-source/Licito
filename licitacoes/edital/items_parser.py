"""Leitura DETERMINÍSTICA dos itens do Anexo I (proposta de preços) de um edital.

Suporta .md (tabela com barras), .docx e .pdf. Não usa IA e não inventa nada:
se não achar a tabela, devolve lista vazia + aviso (a tela oferece tabela manual).
Valores em centavos. Descrições mantidas como no edital (sem corrigir digitação).
"""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path

from licitacoes.services.money import format_brl, parse_brl_to_cents


@dataclass(frozen=True)
class EditalItem:
    number: int
    description: str
    unit: str
    quantity: int
    ceiling_cents: int | None
    position: int


@dataclass
class ParseResult:
    items: list[EditalItem] = field(default_factory=list)
    estimated_total_cents: int | None = None
    computed_total_cents: int | None = None
    method: str = "deterministic"
    needs_review: bool = False
    warnings: list[str] = field(default_factory=list)
    sha256: str = ""
    source_name: str = ""


# ---------------------------------------------------------------- utilidades
def _clean(text) -> str:
    if text is None:
        return ""
    t = str(text).replace("<br>", " ").replace("\\$", "$")
    t = re.sub(r"</?[a-zA-Z][^>]*>", "", t)          # <sup>, <strong>...
    t = t.replace("**", "")
    return re.sub(r"\s+", " ", t).strip()


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


_EST = re.compile(r"VALOR\s+ESTIMADO.{0,120}?R\\?\$\s*([\d.]+,\d{2})", re.I | re.S)


def find_estimated_total(text: str) -> int | None:
    m = _EST.search(text or "")
    return parse_brl_to_cents(m.group(1)) if m else None


def _is_header(cells: list[str]) -> bool:
    up = [c.upper() for c in cells if c]
    return bool(up) and up[0] == "ITEM" and any(c.startswith("ESPECIFICA") for c in up)


_INT = re.compile(r"^\d+$")


def _parse_item_row(cells: list[str], position: int) -> EditalItem | None:
    ne = [c for c in cells if c]
    if len(ne) < 3 or not _INT.match(ne[0]):
        return None
    number, desc, rest = int(ne[0]), ne[1], ne[2:]
    unit, qty, ceiling = "", None, None
    i = 0
    if rest and not _INT.match(rest[0]) and parse_brl_to_cents(rest[0]) is None:
        unit, i = rest[0], 1
    if i < len(rest) and _INT.match(rest[i]):
        qty, i = int(rest[i]), i + 1
    if i < len(rest):
        ceiling = parse_brl_to_cents(rest[i])
    if qty is None:
        return None
    return EditalItem(number, desc, unit, qty, ceiling, position)


def _rows_to_items(rows: list[list[str]]) -> list[EditalItem]:
    """Fluxo de linhas: acha o cabeçalho ITEM/ESPECIFICAÇÕES e lê até 'VALOR TOTAL'."""
    items: list[EditalItem] = []
    inside = False
    for raw in rows:
        cells = [_clean(c) for c in raw]
        if _is_header(cells):
            inside = True
            continue
        if not inside:
            continue
        first = next((c for c in cells if c), "")
        if first.upper().startswith("VALOR TOTAL"):
            break
        it = _parse_item_row(cells, len(items))
        if it:
            items.append(it)
    return items


def _best(candidates: list[list[EditalItem]]) -> list[EditalItem]:
    """Entre várias tabelas candidatas, prefere a que tem preço teto preenchido."""
    if not candidates:
        return []
    return max(candidates, key=lambda its: (sum(i.ceiling_cents is not None for i in its), len(its)))


# ---------------------------------------------------------------- por formato
def _md_tables(text: str) -> list[list[list[str]]]:
    tables, cur = [], []
    for line in text.splitlines():
        s = line.strip()
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if all(re.fullmatch(r"-{3,}|:?-+:?", c) or c == "" for c in cells) and any(cells):
                continue                                  # linha separadora
            cur.append(cells)
        elif cur:
            tables.append(cur)
            cur = []
    if cur:
        tables.append(cur)
    return tables


def _parse_md(path: Path):
    text = path.read_text(encoding="utf-8", errors="replace")
    items = _best([_rows_to_items(t) for t in _md_tables(text)])
    return items, text


def _parse_docx(path: Path):
    import docx
    from docx.table import _Cell

    d = docx.Document(str(path))
    cands = []
    for t in d.tables:
        rows = [[_Cell(tc, t).text for tc in r._tr.tc_lst] for r in t.rows]
        cands.append(_rows_to_items(rows))
    text = "\n".join(p.text for p in d.paragraphs)
    for t in d.tables:                                    # valor estimado pode estar em tabela
        for r in t.rows:
            text += "\n" + " ".join(_Cell(tc, t).text for tc in r._tr.tc_lst)
    return _best(cands), text


def _parse_pdf(path: Path):
    import pymupdf

    rows, text = [], ""
    with pymupdf.open(str(path)) as doc:
        for page in doc:
            text += page.get_text() + "\n"
            for tab in page.find_tables():                # tabelas seguidas = um só fluxo
                rows.extend([[c or "" for c in r] for r in tab.extract()])
    return _rows_to_items(rows), text


# ---------------------------------------------------------------- API
def parse_edital(path: str | Path) -> ParseResult:
    path = Path(path)
    res = ParseResult(sha256=_sha256(path), source_name=path.name)
    ext = path.suffix.lower()
    try:
        if ext in (".md", ".txt"):
            items, text = _parse_md(path)
        elif ext == ".docx":
            items, text = _parse_docx(path)
        elif ext == ".pdf":
            items, text = _parse_pdf(path)
        else:
            res.warnings.append(f"Formato '{ext}' não suportado. Use .md, .docx ou .pdf.")
            return res
    except Exception as exc:                              # erro de leitura: avisa, não inventa
        res.warnings.append(f"Não consegui ler o arquivo ({type(exc).__name__}: {exc}).")
        return res

    res.items = items
    res.estimated_total_cents = find_estimated_total(text)
    if not items:
        res.warnings.append(
            "Não encontrei a tabela de itens (Anexo I). Confira o arquivo ou preencha os itens manualmente.")
        res.needs_review = True
        return res

    if any(i.ceiling_cents is None for i in items):
        res.needs_review = True
        res.warnings.append("Alguns itens estão sem preço máximo. Confira no edital.")
    else:
        res.computed_total_cents = sum(i.quantity * i.ceiling_cents for i in items)
        if res.estimated_total_cents is not None and res.computed_total_cents != res.estimated_total_cents:
            diff = res.computed_total_cents - res.estimated_total_cents
            res.needs_review = True
            res.warnings.append(
                f"A soma dos itens ({format_brl(res.computed_total_cents)}) difere do valor estimado do "
                f"edital ({format_brl(res.estimated_total_cents)}): diferença de {format_brl(diff)}.")
    nums = [i.number for i in items]
    if len(set(nums)) != len(nums):
        res.needs_review = True
        res.warnings.append("Há números de item repetidos. Confira a lista.")
    return res
